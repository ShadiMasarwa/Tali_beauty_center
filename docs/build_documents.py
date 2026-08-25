from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

ROOT = Path(__file__).resolve().parent
ACCENT = "B8797B"
INK = "362E2C"
MUTED = "776965"

def rtl(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    ppr.append(bidi)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcpr.append(shd)

def setup(doc, title, subtitle):
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(0.8)
    sec.left_margin = sec.right_margin = Inches(0.85)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"; normal.font.size = Pt(11); normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.15
    for style_name, size in [("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)]:
        st = styles[style_name]; st.font.name = "Arial"; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor.from_string(ACCENT)
        st.paragraph_format.space_before = Pt(12); st.paragraph_format.space_after = Pt(6)
    p = doc.add_paragraph(); rtl(p)
    run = p.add_run(title); run.font.name="Arial"; run.font.size=Pt(25); run.font.bold=True; run.font.color.rgb=RGBColor.from_string(INK)
    p = doc.add_paragraph(); rtl(p)
    run = p.add_run(subtitle); run.font.name="Arial"; run.font.size=Pt(11); run.font.color.rgb=RGBColor.from_string(MUTED)
    p = doc.add_paragraph(); rtl(p); p.add_run("TALIÉ BEAUTY STUDIO").bold = True
    footer = sec.footer.paragraphs[0]; rtl(footer); footer.add_run("פרויקט גמר — מערכת לניהול תורים למכון יופי")

def heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}"); rtl(p); return p

def para(doc, text, bold_prefix=None):
    p = doc.add_paragraph(); rtl(p)
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True; p.add_run(text[len(bold_prefix):])
    else: p.add_run(text)
    return p

def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet"); rtl(p); p.add_run(item)

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.RIGHT; t.style="Table Grid"; t.autofit=False
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,"EAD9D5"); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs: rtl(p); p.runs[0].bold=True
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs: rtl(p)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    return t

def create_spec():
    doc=Document(); setup(doc,"מסמך אפיון סופי","מערכת TALIÉ לניהול תורים במכון היופי של טלי")
    heading(doc,"1. מטרת המערכת")
    para(doc,"מערכת אינטרנטית בעברית המאפשרת ללקוחות לקבוע תורים ללא הרשמה, ולמנהלת ולעובדות לנהל יומן יחיד למכון. המערכת מונעת קביעת שני תורים במקביל ומתחשבת במשך הטיפולים, שעות העבודה, חסימות וחמש דקות מרווח לכל טיפול.")
    heading(doc,"2. קהלי יעד")
    table(doc,["קהל","שימוש עיקרי"],[["לקוחה","בחירת טיפולים, צפייה בזמינות וקביעת תור"],["עובדת","צפייה ביומן, הוספה ועדכון של תורים"],["מנהלת","כל פעולות העובדת וכן ניהול עובדים, שעות פעילות וחסימות"]],[1.7,4.8])
    heading(doc,"3. דרישות פונקציונליות")
    bullets(doc,["בחירת טיפול אחד או יותר והצגת משך הזמן הכולל.","דרישת שם ומספר טלפון תקין; הערות הן אופציונליות.","הצעת שעות פנויות בלבד, בקפיצות של 15 דקות.","הוספת חמש דקות מרווח לכל טיפול בסדרה.","יומן יום ושבוע לפי שעון ישראל, עם צבע אחיד לפי הטיפול הראשון.","עדכון בזמן אמת באמצעות Socket.IO.","ניהול סטטוסים: נקבע, הסתיים, לא הופיעה ובוטל.","תיעוד מקור התור: אתר או ידני.","דשבורד חודשי לפי סטטוס, סוג טיפול ומקור.","ייצוא תורים ולקוחות לקובצי CSV.","יומן פעילות המתעד פעולות צוות."])
    heading(doc,"4. כללי מספרי טלפון")
    bullets(doc,["ניתן להזין ספרות ומקפים בלבד.","מספר המתחיל ב־05 כולל 10 ספרות; מספר אחר כולל 9 ספרות.","במסד הנתונים נשמרות ספרות בלבד.","בהצגה נעשה שימוש בתבניות 05X-XXX-XXXX או 0X-XXX-XXXX."])
    heading(doc,"5. טכנולוגיות")
    table(doc,["שכבה","טכנולוגיה"],[["לקוח","React, Vite, Bootstrap RTL"],["שרת","Node.js, Express"],["מסד נתונים","MongoDB מקומי, Mongoose"],["אבטחה","Bcrypt, Session, Helmet, Rate Limit"],["זמן אמת","Socket.IO"],["בדיקות","Node Test Runner"]],[1.7,4.8])
    heading(doc,"6. מבנה מסד הנתונים")
    table(doc,["אוסף","שדות עיקריים","קשרים"],[["users","name, username, passwordHash, role","יצירת תורים ויומן פעילות"],["services","code, name, durationMinutes, color","משויך לתורים"],["appointments","customerName, phone, serviceIds, date, times, status, source","שירותים ומשתמש יוצר"],["settings","workingHours, closures","מסמך הגדרות יחיד"],["auditlogs","userName, action, details, createdAt","משתמש מבצע"]],[1.3,3.2,2.0])
    heading(doc,"7. זרימת קביעת תור")
    table(doc,["שלב","פעולה"],[[1,"קליטת שם, טלפון, הערות וטיפולים"],[2,"חישוב משך הטיפולים וחמש דקות לכל טיפול"],[3,"בדיקת שעות פעילות וחסימות"],[4,"בדיקת התנגשות מול כל התורים הקיימים"],[5,"הצגת משבצות פנויות בלבד"],[6,"שמירת התור ושליחת עדכון בזמן אמת"]],[0.8,5.7])
    heading(doc,"8. אבטחה והרשאות")
    bullets(doc,["סיסמאות נשמרות כגיבוב Bcrypt ואינן מוחזרות ללקוח.","עוגיית Session היא HttpOnly ו־SameSite=Lax.","Session נשמר ב־MongoDB למשך 12 שעות.","ניסיונות התחברות מוגבלים בחלון של 15 דקות.","פעולות ניהול רגישות מוגנות בהרשאת מנהלת.","כל קלט מרכזי נבדק בצד השרת.","קובצי .env אינם נכללים במאגר או בחבילת ההגשה."])
    heading(doc,"9. תרחישי קבלה")
    table(doc,["תרחיש","תוצאה צפויה"],[["תור חופף לתור קיים","המערכת דוחה את הפעולה"],["טיפול חורג משעת הסגירה","המועד אינו מוצע"],["חסימה באמצע הטיפול","המועד אינו מוצע"],["עדכון טיפול מאריך את התור","נשמר רק אם אין התנגשות"],["לקוחה קובעת באתר","התור מופיע מיד ביומן הצוות"],["עובדת ניגשת לניהול עובדים","השרת מחזיר שגיאת הרשאה"]],[2.5,4.0])
    doc.save(ROOT/"final-specification.docx")

def create_guide():
    doc=Document(); setup(doc,"מדריך משתמש והפעלה","התקנה, שימוש ותחזוקה של מערכת TALIÉ")
    heading(doc,"1. התקנה")
    for n,text in enumerate(["התקינו Node.js 20 ומעלה ו־MongoDB Community Server.","פתחו Terminal בתיקיית הפרויקט והריצו npm install ולאחר מכן npm run install:all.","העתיקו את קובצי .env.example לקובצי .env ועדכנו SESSION_SECRET.","ודאו ש־MongoDB פועל והריצו npm run seed.","הריצו npm run dev."],1): para(doc,f"{n}. {text}")
    heading(doc,"2. כניסה ראשונית")
    table(doc,["פריט","ערך"],[["כתובת אתר לקוחה","http://localhost:5173"],["כתובת ניהול","http://localhost:5173/admin"],["שם משתמש","tali"],["סיסמה ראשונית","הערך שהוגדר ב־INITIAL_ADMIN_PASSWORD"]],[2.2,4.3])
    para(doc,"יש לשנות את הסיסמה הראשונית לאחר הכניסה הראשונה.")
    heading(doc,"3. שימוש באתר הלקוחה")
    bullets(doc,["מזינים שם מלא ומספר טלפון.","בוחרים טיפול אחד או יותר.","בוחרים תאריך ושעה מתוך השעות הפנויות.","לוחצים על קביעת הטיפול וממתינים לאישור."])
    heading(doc,"4. שימוש במסך הניהול")
    bullets(doc,["יומן תורים: מעבר בין יום לשבוע, ניווט בתאריכים ופתיחת תור לעריכה.","תור חדש: הזנת לקוחה, בחירת טיפולים וקבלת שעות זמינות.","דשבורד: נתוני החודש לפי סטטוס, טיפול ומקור.","לקוחות: חיפוש וייצוא CSV.","טיפולים: עדכון משך טיפול.","יומן פעילות: צפייה בפעולות שבוצעו.","ניהול מנהלת: עובדים, שעות פעילות וחסימות."])
    heading(doc,"5. גיבוי וייצוא")
    para(doc,"במסך הלקוחות ניתן לייצא רשימת לקוחות, ובמסך היומן ניתן לייצא את כל התורים. הקבצים נשמרים בפורמט CSV עם קידוד המתאים לעברית.")
    heading(doc,"6. פתרון תקלות")
    table(doc,["בעיה","בדיקה ופתרון"],[["אין חיבור למסד הנתונים","בדקו ששירות MongoDB פועל ושכתובת MONGODB_URI נכונה"],["הלקוח לא מתחבר לשרת","בדקו את VITE_API_URL ואת CLIENT_URL"],["הכניסה נכשלת","הריצו npm run seed ובדקו את פרטי המשתמש"],["אין שעות פנויות","בדקו שעות עבודה, חסימות ותורים קיימים"],["שינוי בזמן אינו מופיע","בדקו שהשרת ו־Socket.IO פועלים בפורט 5000"]],[2.4,4.1])
    heading(doc,"7. פקודות שימושיות")
    table(doc,["פקודה","מטרה"],[["npm run dev","הפעלת הלקוח והשרת"],["npm run seed","יצירת נתוני אתחול"],["npm test","הרצת בדיקות"],["npm run build","בניית צד הלקוח לייצור"]],[2.2,4.3])
    doc.save(ROOT/"user-guide.docx")

create_spec(); create_guide()
